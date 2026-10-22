#!/usr/bin/env python3
"""Export the dashboard as a slide deck (PPTX + PDF).

Every collapsible visual in the running Streamlit app is captured exactly as the audience sees it (answer text and
chart) and placed on its own slide, one section per divider. Nothing is re-computed here: the deck is the dashboard.

Usage (the dashboard must be running, default http://localhost:8501):
    uv run --group dashboard --with playwright python scripts/export_dashboard_deck.py [--url URL] [--out presentation/retail_dashboard_deck]

Requires Chromium for Playwright. Set DASHBOARD_CHROME to a Chrome/Chromium binary if Playwright's own browsers are not installed.
PDF conversion uses LibreOffice (`soffice`) when available.
"""
from __future__ import annotations

import argparse
import asyncio
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "dashboard"))
import swd  # noqa: E402  (palette only)

SECTIONS = ["Solution map", "Data & model", "Row flow", "Management", "Finance", "AP / Ops", "Sales", "Stores", "Analysis", "Roadmap & AI"]
HIDE_CSS = """
  [data-testid="stSidebar"], [data-testid="stHeader"], header, .modebar { display: none !important; }
  div[data-testid="stDownloadButton"], div[data-testid="stButton"] { display: none !important; }
  section.stMain .block-container { max-width: 1500px; padding-top: 0.5rem; }
"""
W, H = Inches(13.333), Inches(7.5)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#"))


async def capture(url: str, workdir: Path) -> list[dict]:
    """Return [{section, title, png}] for the header tiles, the solution-map diagram and every expander."""
    from playwright.async_api import async_playwright

    exe = os.getenv("DASHBOARD_CHROME")
    shots: list[dict] = []
    async with async_playwright() as p:
        browser = await p.chromium.launch(executable_path=exe, headless=True, args=["--no-sandbox"]) if exe else await p.chromium.launch(headless=True)
        for idx, section in enumerate(SECTIONS, start=1):
            page = await browser.new_page(viewport={"width": 1500, "height": 1100}, device_scale_factor=2)
            await page.goto(f"{url}/?section={idx}", wait_until="networkidle")
            for _ in range(80):
                await page.wait_for_timeout(1500)
                busy = await page.locator("[data-testid='stStatusWidget']").count()
                ready = await page.locator("[data-testid='stExpander']").count()
                if busy == 0 and ready > 0:
                    break
            await page.add_style_tag(content=HIDE_CSS)
            await page.wait_for_timeout(2500)
            # Expand every collapsed expander so the capture includes its body.
            closed = page.locator("[data-testid='stExpander'] details:not([open]) summary")
            for i in range(await closed.count()):
                await closed.nth(0).click()
                await page.wait_for_timeout(600)
            # Titles are read before the headers are hidden; the header would otherwise repeat the slide title.
            expanders = page.locator("[data-testid='stExpander']")
            titles = [(await expanders.nth(i).locator("summary").inner_text()).strip() for i in range(await expanders.count())]
            await page.add_style_tag(content="[data-testid='stExpander'] details > summary { display: none !important; }")
            await page.wait_for_timeout(1500)
            if idx == 1:
                tiles = page.locator("div[data-testid='stHorizontalBlock']").first
                path = workdir / "tiles.png"
                await tiles.screenshot(path=str(path))
                shots.append({"section": "Headline", "title": "One tile per stakeholder note", "png": path})
                diagram = page.locator("div.stMarkdown", has_text="From one CSV to four deliverables").first
                blocks = page.locator("section.stMain div[data-testid='stVerticalBlock'] > div")
                # Capture the region from the diagram heading to the caption below the output boxes.
                heading_box = await diagram.bounding_box()
                caption = page.locator("div[data-testid='stCaptionContainer']", has_text="Vendor marks").first
                cap_box = await caption.bounding_box()
                if heading_box and cap_box:
                    clip = {"x": heading_box["x"] - 8, "y": heading_box["y"] - 4, "width": heading_box["width"] + 16, "height": cap_box["y"] + cap_box["height"] - heading_box["y"] + 8}
                    path = workdir / "solution_map.png"
                    await page.screenshot(path=str(path), clip=clip)
                    shots.append({"section": section, "title": "From one CSV to four deliverables", "png": path})
            for i in range(await expanders.count()):
                exp = expanders.nth(i)
                await exp.scroll_into_view_if_needed()
                await page.wait_for_timeout(500)
                title = titles[i]
                path = workdir / f"s{idx:02d}_{i:02d}.png"
                await exp.screenshot(path=str(path))
                shots.append({"section": section, "title": title, "png": path})
            await page.close()
        await browser.close()
    return shots


def add_text(slide, text, x, y, w, h, size=14, bold=False, color=swd.TEXT, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Open Sans"
    run.font.color.rgb = rgb(color)
    return tb


def add_footer(slide, text):
    add_text(slide, text, Inches(0.5), H - Inches(0.45), W - Inches(1.0), Inches(0.3), size=9, color=swd.GREY_DARK)


def picture_fit(slide, png: Path, x, y, max_w, max_h):
    with Image.open(png) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    pic_w, pic_h = int(w * scale), int(h * scale)
    slide.shapes.add_picture(str(png), x + int((max_w - pic_w) / 2), y, width=Emu(pic_w), height=Emu(pic_h))


def build_deck(shots: list[dict], out_pptx: Path, url: str) -> int:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # title slide
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid(); bg.fill.fore_color.rgb = rgb(swd.TEXT); bg.line.fill.background()
    bar = s.shapes.add_shape(1, 0, 0, Inches(0.18), H)
    bar.fill.solid(); bar.fill.fore_color.rgb = rgb(swd.ACCENT); bar.line.fill.background()
    add_text(s, "RETAIL PERFORMANCE  /  DASHBOARD DECK", Inches(0.85), Inches(0.8), Inches(11), Inches(0.4), size=12, bold=True, color=swd.ACCENT)
    add_text(s, "What the data is, how it flows,\nand what is working", Inches(0.85), Inches(1.4), Inches(11), Inches(1.8), size=34, bold=True, color="FFFFFF")
    add_text(s, "Every slide is a visual from the Streamlit dashboard, captured as the audience sees it.\n"
                "Profit means gross profit on standard cost (pack-adjusted), not net profit.", Inches(0.85), Inches(3.5), Inches(11), Inches(1.0), size=14, color="D3D8DD")
    add_text(s, f"Source: dbt marts in retail.duckdb  |  App: {url}", Inches(0.85), H - Inches(0.9), Inches(11), Inches(0.4), size=11, color="B8C0C7")

    # agenda
    s = prs.slides.add_slide(blank)
    add_text(s, "Agenda: one tab per audience", Inches(0.6), Inches(0.4), Inches(12), Inches(0.6), size=26, bold=True)
    counts = {sec: sum(1 for x in shots if x["section"] == sec) for sec in SECTIONS}
    y = Inches(1.3)
    for i, sec in enumerate(SECTIONS, start=1):
        add_text(s, f"{i:02d}", Inches(0.8), y, Inches(0.7), Inches(0.45), size=18, bold=True, color=swd.ACCENT_DARK)
        add_text(s, sec, Inches(1.5), y, Inches(5), Inches(0.45), size=16, bold=True)
        add_text(s, f"{counts[sec]} visual{'s' if counts[sec] != 1 else ''}", Inches(6.5), y, Inches(3), Inches(0.45), size=12, color=swd.GREY_DARK)
        y += Inches(0.52)
    add_footer(s, "Generated from the dashboard; re-run scripts/export_dashboard_deck.py after any change.")

    current = None
    n = 0
    for shot in shots:
        if shot["section"] != current and shot["section"] != "Headline":
            current = shot["section"]
            s = prs.slides.add_slide(blank)
            band = s.shapes.add_shape(1, 0, Inches(3.0), W, Inches(1.5))
            band.fill.solid(); band.fill.fore_color.rgb = rgb(swd.BG_SOFT); band.line.fill.background()
            idx = SECTIONS.index(current) + 1
            add_text(s, f"{idx:02d}", Inches(0.8), Inches(3.15), Inches(1.2), Inches(1.2), size=40, bold=True, color=swd.ACCENT_DARK)
            add_text(s, current, Inches(2.0), Inches(3.3), Inches(10), Inches(1.0), size=32, bold=True)
        s = prs.slides.add_slide(blank)
        title = re.sub(r"\s+", " ", shot["title"].replace("keyboard_arrow_down", "").replace("keyboard_arrow_right", "")).strip()
        add_text(s, title, Inches(0.5), Inches(0.3), W - Inches(1.0), Inches(0.9), size=20, bold=True)
        picture_fit(s, shot["png"], Inches(0.5), Inches(1.25), W - Inches(1.0), H - Inches(1.85))
        add_footer(s, f"{shot['section']}  |  dashboard capture  |  gross profit, not net profit")
        n += 1
    prs.save(out_pptx)
    return n


def to_pdf(pptx: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found; PDF skipped.")
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return pptx.with_suffix(".pdf")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--url", default=os.getenv("DASHBOARD_URL", "http://localhost:8501"))
    ap.add_argument("--out", default=str(ROOT / "presentation" / "retail_dashboard_deck"))
    args = ap.parse_args()
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        shots = asyncio.run(capture(args.url.rstrip("/"), Path(tmp)))
        n = build_deck(shots, out.with_suffix(".pptx"), args.url)
    print(f"Wrote {out.with_suffix('.pptx')} ({n} visual slides)")
    pdf = to_pdf(out.with_suffix(".pptx"))
    if pdf and pdf.exists():
        print(f"Wrote {pdf}")


if __name__ == "__main__":
    main()
